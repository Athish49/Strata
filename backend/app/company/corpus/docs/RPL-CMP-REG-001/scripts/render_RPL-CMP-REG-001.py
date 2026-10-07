#!/usr/bin/env python3
"""
render_RPL-CMP-REG-001.py

Renders RPL-CMP-REG-001_v4.0 into DOCX, PDF, and XLSX formats.

Usage:
    python3 render_RPL-CMP-REG-001.py

Output files are written to:
    ../render/RPL-CMP-REG-001_v4.0.docx
    ../render/RPL-CMP-REG-001_v4.0.pdf
    ../render/RPL-CMP-REG-001_v4.0.xlsx

Requirements:
    python-docx, reportlab, openpyxl, pandas

Python venv: /Users/athish/Documents/Strata/backend/.venv/bin/python3
"""

import os
import sys
import re
import csv
from pathlib import Path
from datetime import datetime

# ─── Path setup ────────────────────────────────────────────────────────────────
SCRIPT_DIR = Path(__file__).parent.resolve()
DOC_DIR = SCRIPT_DIR.parent
RENDER_DIR = DOC_DIR / "render"
DATA_DIR = DOC_DIR / "data"
MD_FILE = DOC_DIR / "RPL-CMP-REG-001_v4.0.md"
CSV_FILE = DATA_DIR / "compliance_register_2024-12-31.csv"
RENDER_DIR.mkdir(exist_ok=True)

DOCX_OUT = RENDER_DIR / "RPL-CMP-REG-001_v4.0.docx"
PDF_OUT = RENDER_DIR / "RPL-CMP-REG-001_v4.0.pdf"
XLSX_OUT = RENDER_DIR / "RPL-CMP-REG-001_v4.0.xlsx"

# ─── Document metadata ──────────────────────────────────────────────────────────
META = {
    "title": "IURC Regulatory Compliance Register",
    "doc_id": "RPL-CMP-REG-001",
    "version": "4.0",
    "company": "Rockridge Power & Light Company",
    "iurc_id": "99012",
    "approved": "2025-01-23",
    "effective": "2025-01-23",
    "owner": "Marcus Lee (P09), Manager Regulatory Compliance",
    "reviewer": "Elena Vasquez (P08), Director Regulatory Affairs",
    "approver": "Jonathan Pierce (P03), VP Legal & Regulatory",
    "law_as_of": "2024-12-31",
    "snapshot": "S1",
}

# ─── Helpers ────────────────────────────────────────────────────────────────────

def strip_clause_ids(text: str) -> str:
    """Remove HTML comment clause ID markers from Markdown text."""
    return re.sub(r'<!--[^>]+-->', '', text).strip()


def read_markdown() -> str:
    """Read and return the canonical Markdown source."""
    with open(MD_FILE, encoding='utf-8') as f:
        return f.read()


def read_csv() -> list[dict]:
    """Read the compliance register CSV and return a list of row dicts."""
    with open(CSV_FILE, encoding='utf-8', newline='') as f:
        return list(csv.DictReader(f))


def parse_sections(md_text: str) -> list[tuple[str, str]]:
    """
    Parse top-level sections from the Markdown.
    Returns list of (heading, body) tuples.
    """
    # Split on ## headers
    parts = re.split(r'\n(#{1,2} .+)\n', md_text)
    sections = []
    if parts:
        # First part before any heading
        if parts[0].strip():
            sections.append(('', parts[0]))
        for i in range(1, len(parts) - 1, 2):
            heading = parts[i].strip()
            body = parts[i + 1] if i + 1 < len(parts) else ''
            sections.append((heading, body))
    return sections


# ─── DOCX Renderer ──────────────────────────────────────────────────────────────

def render_docx():
    """Generate the DOCX output using python-docx."""
    try:
        from docx import Document
        from docx.shared import Pt, Inches, RGBColor
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml.ns import qn
        from docx.oxml import OxmlElement
    except ImportError:
        print("WARNING: python-docx not available. Skipping DOCX render.")
        return

    doc = Document()

    # Page margins
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1.25)
        section.right_margin = Inches(1.25)

    # Styles
    style = doc.styles['Normal']
    style.font.name = 'Calibri'
    style.font.size = Pt(10)

    # Cover block
    title_para = doc.add_heading(META['title'], level=0)
    title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph()
    meta_table = doc.add_table(rows=6, cols=2)
    meta_table.style = 'Table Grid'
    pairs = [
        ('Document ID / Version', f"{META['doc_id']} v{META['version']}"),
        ('Company / IURC ID', f"{META['company']} | {META['iurc_id']}"),
        ('Law as of / Snapshot', f"{META['law_as_of']} / {META['snapshot']}"),
        ('Approved / Effective', f"{META['approved']} / {META['effective']}"),
        ('Owner', META['owner']),
        ('Approver', META['approver']),
    ]
    for i, (k, v) in enumerate(pairs):
        meta_table.cell(i, 0).text = k
        meta_table.cell(i, 1).text = v
        meta_table.cell(i, 0).paragraphs[0].runs[0].bold = True

    doc.add_page_break()

    # Parse Markdown and add sections
    md_text = read_markdown()
    lines = md_text.split('\n')

    i = 0
    in_table = False
    table_lines = []

    while i < len(lines):
        line = strip_clause_ids(lines[i])

        if not line.strip():
            if in_table:
                _add_md_table(doc, table_lines)
                table_lines = []
                in_table = False
            i += 1
            continue

        # Table detection
        if '|' in line and '---' not in line:
            in_table = True
            table_lines.append(line)
            i += 1
            continue
        elif in_table and '---' in line:
            # Skip separator row
            i += 1
            continue
        elif in_table and '|' not in line:
            _add_md_table(doc, table_lines)
            table_lines = []
            in_table = False
            # Don't increment — re-process this line
            continue

        if in_table:
            i += 1
            continue

        # Headings
        if line.startswith('### '):
            doc.add_heading(line[4:].strip(), level=3)
        elif line.startswith('## '):
            doc.add_heading(line[3:].strip(), level=2)
        elif line.startswith('# '):
            doc.add_heading(line[2:].strip(), level=1)
        elif line.startswith('---'):
            doc.add_paragraph('─' * 60)
        elif line.startswith('**') and line.endswith('**') and len(line) < 100:
            p = doc.add_paragraph()
            run = p.add_run(line.strip('*'))
            run.bold = True
        else:
            # Process inline bold and normal text
            para = doc.add_paragraph()
            _add_formatted_run(para, line)

        i += 1

    if in_table and table_lines:
        _add_md_table(doc, table_lines)

    doc.save(str(DOCX_OUT))
    print(f"DOCX written: {DOCX_OUT}")


def _add_formatted_run(para, line: str):
    """Add a paragraph run handling **bold** markers."""
    from docx.shared import Pt
    parts = re.split(r'(\*\*[^*]+\*\*)', line)
    for part in parts:
        if part.startswith('**') and part.endswith('**'):
            run = para.add_run(part[2:-2])
            run.bold = True
        elif part:
            para.add_run(part)


def _add_md_table(doc, lines: list[str]):
    """Parse and add a Markdown table to the document."""
    rows = []
    for line in lines:
        if '---' in line:
            continue
        cells = [c.strip() for c in line.strip('| \n').split('|')]
        if any(c for c in cells):
            rows.append(cells)

    if not rows:
        return

    max_cols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=max_cols)
    table.style = 'Table Grid'

    for r_idx, row in enumerate(rows):
        for c_idx, cell_text in enumerate(row):
            if c_idx < max_cols:
                cell = table.cell(r_idx, c_idx)
                cell.text = cell_text
                if r_idx == 0:
                    for para in cell.paragraphs:
                        for run in para.runs:
                            run.bold = True

    doc.add_paragraph()


# ─── PDF Renderer ───────────────────────────────────────────────────────────────

def render_pdf():
    """Generate the PDF output using reportlab."""
    try:
        from reportlab.lib.pagesizes import letter
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.lib.units import inch
        from reportlab.lib import colors
        from reportlab.platypus import (SimpleDocTemplate, Paragraph, Spacer,
                                        Table, TableStyle, PageBreak, HRFlowable)
        from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_JUSTIFY
    except ImportError:
        print("WARNING: reportlab not available. Skipping PDF render.")
        return

    doc = SimpleDocTemplate(
        str(PDF_OUT),
        pagesize=letter,
        rightMargin=1.25 * inch,
        leftMargin=1.25 * inch,
        topMargin=1.0 * inch,
        bottomMargin=1.0 * inch,
    )

    styles = getSampleStyleSheet()
    base = styles['Normal']

    style_title = ParagraphStyle('DocTitle', parent=base, fontSize=18,
                                  spaceAfter=12, alignment=TA_CENTER, fontName='Helvetica-Bold')
    style_h1 = ParagraphStyle('H1', parent=base, fontSize=13, spaceBefore=14,
                               spaceAfter=6, fontName='Helvetica-Bold', textColor=colors.HexColor('#1a3a5c'))
    style_h2 = ParagraphStyle('H2', parent=base, fontSize=11, spaceBefore=10,
                               spaceAfter=4, fontName='Helvetica-Bold')
    style_body = ParagraphStyle('Body', parent=base, fontSize=9.5,
                                 spaceBefore=3, spaceAfter=3, leading=13, alignment=TA_JUSTIFY)
    style_meta = ParagraphStyle('Meta', parent=base, fontSize=9,
                                 spaceBefore=2, spaceAfter=2, leading=12)

    story = []

    # Title page
    story.append(Spacer(1, 0.5 * inch))
    story.append(Paragraph(META['title'], style_title))
    story.append(Paragraph(f"<b>{META['doc_id']} v{META['version']}</b>", style_title))
    story.append(Spacer(1, 0.3 * inch))

    meta_data = [
        ['Company', f"{META['company']} | IURC ID {META['iurc_id']}"],
        ['Law as of / Snapshot', f"{META['law_as_of']} / {META['snapshot']}"],
        ['Approved / Effective', f"{META['approved']} / {META['effective']}"],
        ['Owner', META['owner']],
        ['Reviewer', META['reviewer']],
        ['Approver', META['approver']],
    ]
    meta_table = Table(meta_data, colWidths=[1.8 * inch, 4.5 * inch])
    meta_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (0, -1), colors.HexColor('#e8edf3')),
        ('FONT', (0, 0), (-1, -1), 'Helvetica', 9),
        ('FONT', (0, 0), (0, -1), 'Helvetica-Bold', 9),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
        ('VALIGN', (0, 0), (-1, -1), 'TOP'),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
    ]))
    story.append(meta_table)
    story.append(PageBreak())

    # Parse and render Markdown content
    md_text = read_markdown()
    lines = md_text.split('\n')

    i = 0
    in_table = False
    table_lines = []

    while i < len(lines):
        line = strip_clause_ids(lines[i])
        stripped = line.strip()

        if not stripped:
            if in_table:
                _render_pdf_table(story, table_lines, style_body)
                table_lines = []
                in_table = False
            else:
                story.append(Spacer(1, 0.08 * inch))
            i += 1
            continue

        # Table detection
        if '|' in stripped and not stripped.startswith('#'):
            if '---' not in stripped:
                in_table = True
                table_lines.append(stripped)
            i += 1
            continue

        if in_table:
            if '|' not in stripped:
                _render_pdf_table(story, table_lines, style_body)
                table_lines = []
                in_table = False
                # Re-process current line
                continue
            i += 1
            continue

        # Clean up markdown bold for reportlab
        rl_line = re.sub(r'\*\*(.+?)\*\*', r'<b>\1</b>', stripped)
        rl_line = re.sub(r'`(.+?)`', r'<font name="Courier">\1</font>', rl_line)

        if stripped.startswith('### '):
            story.append(Paragraph(stripped[4:], style_h2))
        elif stripped.startswith('## '):
            story.append(Paragraph(stripped[3:], style_h1))
        elif stripped.startswith('# '):
            story.append(Paragraph(stripped[2:], style_title))
        elif stripped.startswith('---'):
            story.append(HRFlowable(width='100%', thickness=0.5, color=colors.grey))
            story.append(Spacer(1, 0.05 * inch))
        else:
            try:
                story.append(Paragraph(rl_line, style_body))
            except Exception:
                story.append(Paragraph(stripped, style_body))

        i += 1

    if in_table and table_lines:
        _render_pdf_table(story, table_lines, style_body)

    doc.build(story)
    print(f"PDF written: {PDF_OUT}")


def _render_pdf_table(story, lines: list[str], body_style):
    """Parse Markdown table lines and add as a ReportLab Table."""
    try:
        from reportlab.platypus import Table, TableStyle, Spacer
        from reportlab.lib import colors
        from reportlab.lib.units import inch
    except ImportError:
        return

    rows = []
    for line in lines:
        if '---' in line:
            continue
        cells = [c.strip() for c in line.strip('| \n').split('|')]
        if any(c for c in cells):
            rows.append(cells)

    if not rows:
        return

    max_cols = max(len(r) for r in rows)
    # Pad rows
    padded = [r + [''] * (max_cols - len(r)) for r in rows]

    col_width = 6.0 * inch / max_cols
    col_widths = [col_width] * max_cols

    table = Table(padded, colWidths=col_widths, repeatRows=1)
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1a3a5c')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONT', (0, 0), (-1, 0), 'Helvetica-Bold', 8),
        ('FONT', (0, 1), (-1, -1), 'Helvetica', 7.5),
        ('GRID', (0, 0), (-1, -1), 0.25, colors.HexColor('#cccccc')),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f5f7fa')]),
        ('VALIGN', (0, 0), (-1, -1), 'TOP'),
        ('TOPPADDING', (0, 0), (-1, -1), 3),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
        ('LEFTPADDING', (0, 0), (-1, -1), 4),
        ('RIGHTPADDING', (0, 0), (-1, -1), 4),
        ('WORDWRAP', (0, 0), (-1, -1), True),
    ]))
    story.append(table)
    story.append(Spacer(1, 0.1 * inch))


# ─── XLSX Renderer ──────────────────────────────────────────────────────────────

def render_xlsx():
    """Generate the XLSX output using openpyxl."""
    try:
        import openpyxl
        from openpyxl.styles import (Font, PatternFill, Alignment, Border, Side,
                                      numbers)
        from openpyxl.utils import get_column_letter
    except ImportError:
        print("WARNING: openpyxl not available. Skipping XLSX render.")
        return

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    rows = read_csv()

    # ── Sheet 1: Full Register ──────────────────────────────────────────────────
    ws_reg = wb.create_sheet("Full Register")

    header_fill = PatternFill("solid", fgColor="1A3A5C")
    header_font = Font(bold=True, color="FFFFFF", size=9)
    body_font = Font(size=9)
    alt_fill = PatternFill("solid", fgColor="F0F4F8")
    p1_fill = PatternFill("solid", fgColor="FFE0E0")
    border_side = Side(style='thin', color='CCCCCC')
    thin_border = Border(left=border_side, right=border_side,
                         top=border_side, bottom=border_side)
    wrap = Alignment(wrap_text=True, vertical='top')

    headers = list(rows[0].keys()) if rows else []
    col_widths = {
        'obligation_id': 16, 'citation': 18, 'title': 30, 'description': 50,
        'frequency': 14, 'accountable_owner_id': 12, 'implementation_doc': 18,
        'implementation_clause': 24, 'compliance_control': 30, 'monitoring_method': 35,
        'last_verified': 14, 'next_due': 14, 'status': 12, 'priority': 8, 'notes': 40,
    }

    for c_idx, col_name in enumerate(headers, start=1):
        cell = ws_reg.cell(row=1, column=c_idx, value=col_name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center')
        cell.border = thin_border
        ws_reg.column_dimensions[get_column_letter(c_idx)].width = col_widths.get(col_name, 15)

    ws_reg.row_dimensions[1].height = 30
    ws_reg.freeze_panes = 'A2'

    for r_idx, row in enumerate(rows, start=2):
        for c_idx, col_name in enumerate(headers, start=1):
            cell = ws_reg.cell(row=r_idx, column=c_idx, value=row.get(col_name, ''))
            cell.font = body_font
            cell.alignment = wrap
            cell.border = thin_border
            # P1 rows get light red background
            if row.get('priority') == 'P1':
                cell.fill = p1_fill
            elif r_idx % 2 == 0:
                cell.fill = alt_fill

    # Auto filter
    ws_reg.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"

    # ── Sheet 2: Summary by IAC Title ──────────────────────────────────────────
    ws_sum = wb.create_sheet("Summary by IAC Title")

    from collections import Counter
    iac_groups = {
        '170 IAC 1-2': 'Schedule Filing Rules',
        '170 IAC 1-5': 'General Rate Case Procedures',
        '170 IAC 1-6': '30-Day Admin Filing',
        '170 IAC 1-7': 'Municipal Annexation',
        '170 IAC 4-1': 'Electric Service Rules',
        '170 IAC 4-9': 'Vegetation Management',
        '170 IAC 16-1': 'Customer Dispute Process',
        'IC ': 'Indiana Code Statutory',
    }

    def classify(citation):
        for prefix, name in iac_groups.items():
            if citation.startswith(prefix):
                return name
        return 'Other'

    summary = Counter(classify(r['citation']) for r in rows)

    sum_headers = ['IAC Area', 'Obligation Count', 'P1 Count', 'P2 Count', 'P3 Count']
    for c_idx, h in enumerate(sum_headers, start=1):
        cell = ws_sum.cell(row=1, column=c_idx, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = thin_border

    ws_sum.row_dimensions[1].height = 25

    group_rows = {}
    for row in rows:
        g = classify(row['citation'])
        if g not in group_rows:
            group_rows[g] = {'total': 0, 'P1': 0, 'P2': 0, 'P3': 0}
        group_rows[g]['total'] += 1
        p = row.get('priority', 'P3')
        if p in group_rows[g]:
            group_rows[g][p] += 1

    for r_idx, (group, counts) in enumerate(group_rows.items(), start=2):
        ws_sum.cell(row=r_idx, column=1, value=group).border = thin_border
        ws_sum.cell(row=r_idx, column=2, value=counts['total']).border = thin_border
        ws_sum.cell(row=r_idx, column=3, value=counts['P1']).border = thin_border
        ws_sum.cell(row=r_idx, column=4, value=counts['P2']).border = thin_border
        ws_sum.cell(row=r_idx, column=5, value=counts['P3']).border = thin_border
        if r_idx % 2 == 0:
            for c in range(1, 6):
                ws_sum.cell(row=r_idx, column=c).fill = alt_fill

    for c_idx, w in enumerate([30, 16, 10, 10, 10], start=1):
        ws_sum.column_dimensions[get_column_letter(c_idx)].width = w

    # ── Sheet 3: P1 Obligations ─────────────────────────────────────────────────
    ws_p1 = wb.create_sheet("P1 Obligations")

    for c_idx, col_name in enumerate(headers, start=1):
        cell = ws_p1.cell(row=1, column=c_idx, value=col_name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center')
        cell.border = thin_border
        ws_p1.column_dimensions[get_column_letter(c_idx)].width = col_widths.get(col_name, 15)

    ws_p1.row_dimensions[1].height = 30
    ws_p1.freeze_panes = 'A2'

    p1_rows = [r for r in rows if r.get('priority') == 'P1']
    for r_idx, row in enumerate(p1_rows, start=2):
        for c_idx, col_name in enumerate(headers, start=1):
            cell = ws_p1.cell(row=r_idx, column=c_idx, value=row.get(col_name, ''))
            cell.font = body_font
            cell.alignment = wrap
            cell.border = thin_border
            if r_idx % 2 == 0:
                cell.fill = alt_fill

    ws_p1.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"

    # ── Sheet 4: Cover ─────────────────────────────────────────────────────────
    ws_cover = wb.create_sheet("Cover", 0)
    ws_cover.column_dimensions['A'].width = 30
    ws_cover.column_dimensions['B'].width = 50

    cover_data = [
        ('Document', META['title']),
        ('Document ID', META['doc_id']),
        ('Version', META['version']),
        ('Company', META['company']),
        ('IURC Utility ID', META['iurc_id']),
        ('Law as of Date', META['law_as_of']),
        ('Snapshot', META['snapshot']),
        ('Approved', META['approved']),
        ('Effective', META['effective']),
        ('Owner', META['owner']),
        ('Reviewer', META['reviewer']),
        ('Approver', META['approver']),
        ('Total Obligations', str(len(rows))),
        ('P1 Obligations', str(len([r for r in rows if r.get('priority') == 'P1']))),
        ('P2 Obligations', str(len([r for r in rows if r.get('priority') == 'P2']))),
        ('P3 Obligations', str(len([r for r in rows if r.get('priority') == 'P3']))),
        ('Generated', datetime.now().strftime('%Y-%m-%d %H:%M')),
    ]

    title_cell = ws_cover.cell(row=1, column=1, value=META['title'])
    title_cell.font = Font(bold=True, size=14, color='1A3A5C')
    ws_cover.merge_cells('A1:B1')
    ws_cover.row_dimensions[1].height = 30

    for r_idx, (k, v) in enumerate(cover_data, start=2):
        kc = ws_cover.cell(row=r_idx, column=1, value=k)
        vc = ws_cover.cell(row=r_idx, column=2, value=v)
        kc.font = Font(bold=True, size=10)
        vc.font = Font(size=10)
        kc.border = thin_border
        vc.border = thin_border
        if r_idx % 2 == 0:
            kc.fill = alt_fill
            vc.fill = alt_fill

    wb.save(str(XLSX_OUT))
    print(f"XLSX written: {XLSX_OUT}")


# ─── Main ───────────────────────────────────────────────────────────────────────

def main():
    print(f"Rendering RPL-CMP-REG-001 v4.0 ...")
    print(f"  Source MD:  {MD_FILE}")
    print(f"  Source CSV: {CSV_FILE}")
    print(f"  Output dir: {RENDER_DIR}")
    print()

    render_docx()
    render_pdf()
    render_xlsx()

    print()
    print("Done. Output files:")
    for f in [DOCX_OUT, PDF_OUT, XLSX_OUT]:
        if f.exists():
            size_kb = f.stat().st_size / 1024
            print(f"  {f.name:45s}  {size_kb:8.1f} KB")
        else:
            print(f"  {f.name:45s}  NOT CREATED")


if __name__ == '__main__':
    main()
